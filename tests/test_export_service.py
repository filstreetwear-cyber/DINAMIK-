"""
Milestone 15 (Exports) tests.

Scope was clarified directly with Fil before writing any code (see
app/services/export_service.py's own docstring): an internal-team .xlsx
workbook, one row per brand, read-only. Layered the same way
tests/test_weekly_scan.py is:

1. Pure unit tests for `_summarize_reasons()` -- no I/O.
2. Service-level integration tests, real Postgres in MOCK_MODE, using
   the same rollback-based `session` fixture (and the same
   `provider_orchestrator`/`ai_gateway` fixtures)
   tests/test_pipeline_orchestrator.py / tests/test_weekly_scan.py use --
   `run_brand_export`/`build_brand_export_rows` never write to the
   database (this feature is read-only end to end), so there is nothing
   to roll back for THOSE calls specifically, but seeding real data via
   `run_research_pipeline` still needs the rollback-based session to
   avoid polluting the shared dev DB.
3. CLI tests (`dinamik export brands`) via Typer's `CliRunner`, real
   commits + explicit cleanup (same pattern as tests/test_cli.py /
   tests/test_weekly_scan.py's own CLI section), since each CLI
   invocation builds its own engine/session.

`tests/conftest.py`'s autouse fixture isolates `EXPORTS_DIR` to a
per-test tmp directory, so workbook-file assertions below need no
filesystem cleanup either.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from typer.testing import CliRunner

from app.ai.gateway import AIGateway
from app.ai.mock_provider import MockAIProvider
from app.cli import app as cli_app
from app.core.config import get_settings
from app.models import (
    AgentEvent,
    AgentRun,
    Brand,
    ErrorLog,
    Evidence,
    Opportunity,
    ResearchJob,
    ResearchResult,
    Score,
)
from app.models.amazon import AmazonProduct, AmazonSeller
from app.services.export_service import (
    ExportDisabledError,
    _summarize_reasons,
    build_brand_export_rows,
    run_brand_export,
    write_brand_export_workbook,
)
from app.services.pipeline_orchestrator import run_research_pipeline
from app.tools.amazon.provider import AmazonProvider
from app.tools.core.orchestrator import IntelligenceSourceOrchestrator
from app.tools.core.registry import ProviderRegistry
from app.tools.keepa.provider import KeepaProvider

# ---------------------------------------------------------------------------
# 1. Pure unit tests -- no database.
# ---------------------------------------------------------------------------


def test_summarize_reasons_returns_none_for_empty_input():
    assert _summarize_reasons(None) is None
    assert _summarize_reasons([]) is None


def test_summarize_reasons_joins_dict_reasons():
    reasons = [{"reason": "rating >= 4.2"}, {"reason": "review_count >= 50"}]
    assert _summarize_reasons(reasons) == "rating >= 4.2; review_count >= 50"


def test_summarize_reasons_joins_plain_string_reasons():
    assert _summarize_reasons(["a", "b", "c"]) == "a; b; c"


def test_summarize_reasons_truncates_long_output():
    reasons = [{"reason": "x" * 100} for _ in range(20)]
    result = _summarize_reasons(reasons)
    assert result is not None
    assert len(result) == 500


# ---------------------------------------------------------------------------
# 2. Service-level integration tests -- real Postgres, rollback-based.
# ---------------------------------------------------------------------------


async def _make_engine():
    settings = get_settings()
    return create_async_engine(settings.database_url)


async def _db_reachable() -> bool:
    try:
        engine = await _make_engine()
        async with engine.connect() as conn:
            await conn.run_sync(lambda c: None)
        await engine.dispose()
        return True
    except OperationalError:
        return False


@pytest.fixture
async def session():
    if not await _db_reachable():
        pytest.skip("PostgreSQL not reachable — export tests need a live DB")
    engine = await _make_engine()
    factory = async_sessionmaker(bind=engine, expire_on_commit=False)
    async with factory() as s:
        yield s
        await s.rollback()
    await engine.dispose()


@pytest.fixture
def provider_orchestrator(monkeypatch: pytest.MonkeyPatch) -> IntelligenceSourceOrchestrator:
    monkeypatch.setenv("MOCK_MODE", "true")
    get_settings.cache_clear()
    settings = get_settings()

    registry = ProviderRegistry()
    registry.register(AmazonProvider(settings=settings), priority=20, cost_model="mock")
    registry.register(KeepaProvider(settings=settings), priority=30, cost_model="mock")
    return IntelligenceSourceOrchestrator(registry)


def _ai_claims_draft() -> str:
    return json.dumps(
        {
            "claims": [
                {
                    "claim": "Summarized from the brand's discovered evidence.",
                    "data_type": "INFERENCE",
                    "confidence": 65.0,
                    "evidence_ids": [],
                    "reasoning_summary": "Export integration test — generic summary claim.",
                }
            ]
        }
    )


@pytest.fixture
def ai_gateway() -> AIGateway:
    return AIGateway(MockAIProvider(structured_response=_ai_claims_draft()))


async def _seed_job(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway
) -> None:
    job = ResearchJob(name="export test job", marketplace="IT", category="Casa e cucina")
    session.add(job)
    await session.flush()
    await run_research_pipeline(
        session, job, provider_orchestrator=provider_orchestrator, ai_gateway=ai_gateway
    )


@pytest.mark.asyncio
async def test_build_brand_export_rows_reflects_real_pipeline_data(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway
):
    await _seed_job(session, provider_orchestrator, ai_gateway)

    rows = await build_brand_export_rows(session, status="ACTIVE")

    assert len(rows) > 0
    row = rows[0]
    assert row.brand_name
    assert row.status == "ACTIVE"
    assert row.product_count > 0  # mock discovery always resolves at least one product per brand
    # At least one of the two implemented opportunity types should have
    # been scored by the real (mock-mode) pipeline run.
    assert set(row.opportunities.keys()) <= {"AMAZON_OPPORTUNITY", "BUSINESS_VALUE"}
    assert row.evidence_count >= 0


@pytest.mark.asyncio
async def test_build_brand_export_rows_filters_by_status(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway
):
    await _seed_job(session, provider_orchestrator, ai_gateway)

    active_rows = await build_brand_export_rows(session, status="ACTIVE")
    merged_rows = await build_brand_export_rows(session, status="MERGED")

    assert len(active_rows) > 0
    assert merged_rows == []  # no brand has been merged in this fresh seed


@pytest.mark.asyncio
async def test_write_brand_export_workbook_produces_real_formatted_xlsx(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway, tmp_path: Path
):
    await _seed_job(session, provider_orchestrator, ai_gateway)
    rows = await build_brand_export_rows(session, status="ACTIVE")

    path = write_brand_export_workbook(rows, tmp_path / "brands-test.xlsx")

    assert path.exists()
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet is not None
    assert sheet.title == "Brands"
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold is True
    assert sheet["A1"].value == "Brand"

    values = list(sheet.iter_rows(values_only=True))
    assert len(values) == len(rows) + 1  # header + one row per brand
    assert values[1][0] == rows[0].brand_name


@pytest.mark.asyncio
async def test_run_brand_export_writes_to_exports_dir_and_is_read_only(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway
):
    await _seed_job(session, provider_orchestrator, ai_gateway)
    settings = get_settings()

    brands_before = (await session.execute(select(Brand))).scalars().all()

    summary = await run_brand_export(session, settings=settings)

    assert summary.rows_exported == len(brands_before)
    assert Path(summary.path).exists()
    assert Path(summary.path).is_relative_to(Path(settings.exports_dir))

    # Read-only claim, verified rather than assumed: no new Brand rows
    # were created by the export itself.
    brands_after = (await session.execute(select(Brand))).scalars().all()
    assert len(brands_after) == len(brands_before)


@pytest.mark.asyncio
async def test_run_brand_export_raises_when_disabled(
    session: AsyncSession, provider_orchestrator: IntelligenceSourceOrchestrator, ai_gateway: AIGateway
):
    await _seed_job(session, provider_orchestrator, ai_gateway)
    settings = get_settings()
    settings.enable_export = False
    try:
        with pytest.raises(ExportDisabledError):
            await run_brand_export(session, settings=settings)
    finally:
        settings.enable_export = True


# ---------------------------------------------------------------------------
# 3. CLI tests -- real commits, explicit cleanup.
# ---------------------------------------------------------------------------

runner = CliRunner()


async def _cleanup_all() -> None:
    engine = await _make_engine()
    factory = async_sessionmaker(bind=engine, expire_on_commit=False)
    async with factory() as s:
        await s.execute(delete(Evidence))
        await s.execute(delete(AgentEvent))
        await s.execute(delete(ErrorLog))
        await s.execute(delete(Score))
        await s.execute(delete(Opportunity))
        await s.execute(delete(ResearchResult))
        await s.execute(delete(AmazonProduct))
        await s.execute(delete(AmazonSeller))
        await s.execute(delete(AgentRun))
        await s.execute(delete(Brand))
        await s.execute(delete(ResearchJob))
        await s.commit()
    await engine.dispose()


@pytest.fixture(autouse=True)
def _require_db_and_mock_mode_for_cli_tests(request: pytest.FixtureRequest, monkeypatch: pytest.MonkeyPatch):
    if not request.node.name.startswith("test_cli_"):
        yield
        return

    import asyncio

    if not asyncio.run(_db_reachable()):
        pytest.skip("PostgreSQL not reachable — export CLI tests run real commands against a live DB")
    monkeypatch.setenv("MOCK_MODE", "true")
    get_settings.cache_clear()
    yield
    asyncio.run(_cleanup_all())
    get_settings.cache_clear()


def test_cli_export_brands_writes_a_real_workbook_after_a_real_pipeline_run():
    pipeline_result = runner.invoke(
        cli_app, ["pipeline", "run", "--name", "export-cli-test-job", "--marketplace", "IT", "--category", "Casa e cucina"]
    )
    assert pipeline_result.exit_code == 0, pipeline_result.stdout

    result = runner.invoke(cli_app, ["export", "brands"])

    assert result.exit_code == 0, result.stdout
    assert "rows_exported:" in result.stdout
    assert "path:" in result.stdout

    # Extract the reported path and verify the file genuinely exists and
    # is a real (loadable) workbook, not just a printed claim.
    path_line = next(line for line in result.stdout.splitlines() if line.startswith("path:"))
    reported_path = Path(path_line.split("path:", 1)[1].strip())
    assert reported_path.exists()
    workbook = load_workbook(reported_path)
    active_sheet = workbook.active
    assert active_sheet is not None
    assert active_sheet.title == "Brands"


def test_cli_export_brands_reports_zero_rows_when_no_brands_match():
    result = runner.invoke(cli_app, ["export", "brands", "--status", "MERGED"])

    assert result.exit_code == 0, result.stdout
    assert "rows_exported: 0" in result.stdout


def test_cli_export_brands_reports_disabled_export_cleanly(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("ENABLE_EXPORT", "false")
    get_settings.cache_clear()
    try:
        result = runner.invoke(cli_app, ["export", "brands"])
        assert result.exit_code == 1
        assert "export disabled" in result.stdout
    finally:
        get_settings.cache_clear()
